"""CSV and XLSX readers: one block per non-empty cell, located by sheet, row and column.
Formulas are never evaluated; cached values are read and the cells are flagged."""

import csv
import io
import warnings as python_warnings
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from sequence_vault.adapters.parsers.decoding import decode
from sequence_vault.adapters.parsers.document import DocumentBuilder
from sequence_vault.application.ports import ParseFailed

CSV_VERSION = "csv-1"
XLSX_VERSION = "xlsx-1"


def _cell(sheet: str, row: int, column: int) -> dict[str, Any]:
    return {
        "kind": "spreadsheet_cell",
        "worksheet": sheet,
        "cell": f"{get_column_letter(column + 1)}{row + 1}",
        "row": row,
        "column": column,
    }


def parse_csv(
    data: bytes, *, file_id: str, run_id: str, parse_options: dict[str, Any]
) -> dict[str, Any]:
    text, encoding = decode(data)
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(
            text[:4096], delimiters=",\t;"
        )
    except csv.Error:
        dialect = csv.excel
    builder = DocumentBuilder(file_id, run_id, data, CSV_VERSION, encoding, parse_options)
    for row_index, row in enumerate(csv.reader(io.StringIO(text), dialect)):
        for column, value in enumerate(row):
            if value.strip():
                builder.add(
                    f"s0r{row_index}c{column}", "table_cell", value, _cell("csv", row_index, column)
                )
    return builder.build()


def parse_xlsx(
    data: bytes, *, file_id: str, run_id: str, parse_options: dict[str, Any]
) -> dict[str, Any]:
    try:
        with python_warnings.catch_warnings():
            python_warnings.simplefilter("ignore")
            formulas = load_workbook(io.BytesIO(data), data_only=False, keep_vba=False)
            values = load_workbook(io.BytesIO(data), data_only=True, keep_vba=False)
    except Exception as error:  # openpyxl raises many types for damaged files
        raise ParseFailed(
            "The workbook is damaged or unsupported.", code="corrupt_document"
        ) from error
    builder = DocumentBuilder(file_id, run_id, data, XLSX_VERSION, "utf-8", parse_options)
    formula_cells = hidden_rows = hidden_columns = merged = 0
    hidden_sheets: list[str] = []
    for sheet_index, sheet in enumerate(formulas.worksheets):
        cached = values[sheet.title]
        if sheet.sheet_state != "visible":
            hidden_sheets.append(sheet.title)
        hidden_rows += sum(1 for dim in sheet.row_dimensions.values() if dim.hidden)
        hidden_columns += sum(1 for dim in sheet.column_dimensions.values() if dim.hidden)
        merged += len(sheet.merged_cells.ranges)
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                row_index, column = cell.row - 1, cell.column - 1
                block_id = f"s{sheet_index}r{row_index}c{column}"
                is_formula = cell.data_type == "f"
                value = cell.value
                if is_formula:
                    # Never evaluated: the cached result when Excel stored one, otherwise the
                    # formula text itself, flagged for review either way.
                    stored = cached.cell(row=cell.row, column=cell.column).value
                    value = stored if stored is not None else cell.value
                if value is None or not str(value).strip():
                    continue
                builder.add(
                    block_id, "table_cell", str(value), _cell(sheet.title, row_index, column)
                )
                if is_formula:
                    formula_cells += 1
                    builder.unresolved.append(block_id)
    if formula_cells:
        builder.warnings.append(
            f"{formula_cells} formula cell(s) were not evaluated; cached values need review."
        )
    if hidden_sheets:
        builder.warnings.append(f"Hidden worksheets were read: {', '.join(hidden_sheets)}.")
    if hidden_rows or hidden_columns:
        builder.warnings.append(
            f"Hidden rows ({hidden_rows}) and columns ({hidden_columns}) were read."
        )
    if merged:
        builder.warnings.append(f"{merged} merged cell range(s); check row and column pairing.")
    return builder.build()
